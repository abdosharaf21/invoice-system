"""Unit tests for the reconciliation CSV/XLSX export helpers."""

import csv
import io
from datetime import date, datetime
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from backend.modules.reconciliation.export import (
    CSV_MIMETYPE,
    XLSX_MIMETYPE,
    build_file,
    render_csv,
    render_xlsx,
)

_RESULT_ROW = {
    "id": 1,
    "run_id": 7,
    "match_status": "mismatched",
    "account_invoice_id": 10,
    "account_invoice_number": "INV-100",
    "account_uuid": "d3c6e4f7-1a2b-4c3d-8e5f-6a7b8c9d0e1f",
    "account_invoice_date": date(2024, 3, 1),
    "account_currency": "EGP",
    "counterparty_name": "Acme, Corp. \"HQ\"",
    "counterparty_tax_id": "12345",
    "accounting_subtotal": Decimal("1150.00"),
    "accounting_vat": Decimal("150.00"),
    "accounting_total": Decimal("1300.00"),
    "tax_invoice_id": 20,
    "tax_internal_id": "T-100",
    "tax_uuid": "d3c6e4f7-1a2b-4c3d-8e5f-6a7b8c9d0e1f",
    "tax_issue_datetime": datetime(2024, 3, 1, 9, 30, 0),
    "tax_currency": "EGP",
    "tax_total_sales": Decimal("1150.00"),
    "tax_net_amount": Decimal("1000.00"),
    "tax_vat_amount": Decimal("150.00"),
    "tax_total_amount": Decimal("1150.00"),
    "discrepancy_amount": Decimal("150.00"),
    "notes": "line one\n\"quoted\", comma",
    "created_at": datetime(2024, 4, 1, 10, 0, 0),
}

_ERROR_ROW = {
    "id": 2,
    "run_id": 7,
    "source_type": "account",
    "entity_id": 10,
    "error_type": "TOTAL_AMOUNT_MISMATCH",
    "field": "total_amount",
    "accounting_value": "1300.00",
    "tax_authority_value": "1150.00",
    "difference": Decimal("150.00"),
    "message": "total_amount differs",
    "created_at": datetime(2024, 4, 1, 10, 0, 1),
}


# ---------------------------------------------------------------------------
# CSV
# ---------------------------------------------------------------------------


def test_render_csv_emits_headers_and_escapes_values():
    text = render_csv(
        (
            ("match_status", "match_status"),
            ("counterparty_name", "counterparty_name"),
            ("notes", "notes"),
            ("accounting_subtotal", "accounting_subtotal"),
            ("account_invoice_date", "account_invoice_date"),
        ),
        [_RESULT_ROW],
    )
    lines = text.splitlines()
    assert lines[0] == "match_status,counterparty_name,notes,accounting_subtotal,account_invoice_date"
    assert 'mismatched,"Acme, Corp. ""HQ""","line one\n"\"quoted"", comma",1150.00,2024-03-01' in text
    # Round-trips losslessly through a csv reader.
    parsed = list(csv.reader(io.StringIO(text)))
    assert parsed[1][0] == "mismatched"
    assert parsed[1][1] == 'Acme, Corp. "HQ"'
    assert parsed[1][2] == 'line one\n"quoted", comma'


def test_render_csv_preserves_decimal_and_empty_cells():
    row = {
        "id": 1, "run_id": 7, "source_type": "tax", "entity_id": None,
        "error_type": "INVALID_UUID", "field": "uuid",
        "accounting_value": None, "tax_authority_value": "bad",
        "difference": None, "message": "", "created_at": None,
    }
    text = render_csv(
        (
            ("entity_id", "entity_id"),
            ("accounting_value", "accounting_value"),
            ("difference", "difference"),
            ("message", "message"),
        ),
        [row],
    )
    assert text.splitlines()[1].split(",") == ["", "", "", ""]


def test_render_csv_empty_dataset_has_only_headers():
    text = render_csv((("id", "id"), ("field", "field")), [])
    assert text == "id,field\r\n"


# ---------------------------------------------------------------------------
# XLSX
# ---------------------------------------------------------------------------


def test_render_xlsx_headers_rows_and_types():
    data = render_xlsx("Results", (
        ("result_id", "id"),
        ("match_status", "match_status"),
        ("accounting_subtotal", "accounting_subtotal"),
        ("discrepancy_amount", "discrepancy_amount"),
        ("account_invoice_date", "account_invoice_date"),
        ("created_at", "created_at"),
        ("notes", "notes"),
    ), [_RESULT_ROW])

    workbook = load_workbook(io.BytesIO(data))
    worksheet = workbook.active
    assert worksheet.title == "Results"
    assert list(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))[0] == (
        "result_id", "match_status", "accounting_subtotal",
        "discrepancy_amount", "account_invoice_date", "created_at", "notes",
    )
    values = next(worksheet.iter_rows(min_row=2, max_row=2, values_only=True))
    assert values[0] == 1
    assert values[3] == pytest.approx(150.00)
    assert values[4].year == 2024 and values[4].month == 3
    assert values[5].year == 2024
    assert values[6] == 'line one\n"quoted", comma'


def test_render_xlsx_empty_dataset_is_valid_with_headers():
    data = render_xlsx("Errors", (("error_id", "id"), ("message", "message")), [])
    workbook = load_workbook(io.BytesIO(data))
    worksheet = workbook.active
    assert worksheet.title == "Errors"
    values = list(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))[0]
    assert values == ("error_id", "message")
    assert worksheet.max_row == 1


def test_render_xlsx_writes_error_rows():
    data = render_xlsx("Errors", (
        ("error_id", "id"),
        ("error_type", "error_type"),
        ("field", "field"),
        ("accounting_value", "accounting_value"),
        ("tax_authority_value", "tax_authority_value"),
        ("difference", "difference"),
        ("created_at", "created_at"),
    ), [_ERROR_ROW])
    worksheet = load_workbook(io.BytesIO(data)).active
    values = next(worksheet.iter_rows(min_row=2, max_row=2, values_only=True))
    assert values[1] == "TOTAL_AMOUNT_MISMATCH"
    assert values[2] == "total_amount"
    assert values[3] == "1300.00"
    assert values[5] == pytest.approx(150.00)
    assert values[6].year == 2024


# ---------------------------------------------------------------------------
# build_file
# ---------------------------------------------------------------------------


def test_build_file_csv_returns_csv_payload_and_mimetype():
    payload, mimetype, extension = build_file("results", [_RESULT_ROW], "csv")
    assert mimetype == CSV_MIMETYPE
    assert extension == "csv"
    assert payload.startswith(b"result_id,run_id,match_status")


def test_build_file_xlsx_returns_xlsx_payload_and_mimetype():
    payload, mimetype, extension = build_file("errors", [_ERROR_ROW], "xlsx")
    assert mimetype == XLSX_MIMETYPE
    assert extension == "xlsx"
    workbook = load_workbook(io.BytesIO(payload))
    assert workbook.active.title == "Errors"
    assert workbook.active.max_row == 2


def test_build_file_rejects_unknown_format():
    with pytest.raises(ValueError):
        build_file("results", [], "pdf")


def test_build_file_rejects_unknown_kind():
    with pytest.raises(ValueError):
        build_file("bogus", [], "csv")